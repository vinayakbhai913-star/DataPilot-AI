from fastapi import FastAPI, UploadFile, File
import csv
import io
from collections import Counter, defaultdict
from datetime import datetime
from openpyxl import load_workbook,Workbook
from fastapi.responses import HTMLResponse, StreamingResponse

app = FastAPI(
    title="DataPilot AI",
    description="AI-powered Business Analytics Platform",
    version="2.0.0"
)
latest_report = {}

# -----------------------------
# HOME
# -----------------------------
@app.get("/")
def home():
    return {
        "app": "DataPilot AI",
        "version": "2.0.0",
        "status": "running",
        "message": "Business Analytics Engine is ready 🚀"
    }


# -----------------------------
# HEALTH CHECK
# -----------------------------
@app.get("/health")
def health():
    return {
        "status": "healthy",
        "engine": "DataPilot Analytics Engine"
    }


# -----------------------------
# HELPER FUNCTIONS
# -----------------------------
def to_float(value):
    try:
        return float(str(value).replace(",", "").strip())
    except:
        return 0.0


def to_int(value):
    try:
        return int(float(str(value).strip()))
    except:
        return 0


# -----------------------------
# ANALYTICS ENGINE
# -----------------------------
@app.post("/analyze")
async def analyze(file: UploadFile = File(...)):
    # -----------------------------
    # CHECK FILE TYPE
    # -----------------------------

    filename = file.filename.lower()

    if not filename.endswith((".csv", ".xlsx")):
        return {
            "error": "Only CSV and XLSX files are supported."
        }

    content = await file.read()

    try:
        # Excel XLSX
        if filename.endswith(".xlsx"):
            workbook = load_workbook(
                filename=io.BytesIO(content),
                data_only=True
            )

            sheet = workbook.active
            data = list(sheet.values)

            if not data:
                return {"error": "Excel file is empty."}

            columns = [
                str(column) if column is not None else ""
                for column in data[0]
            ]

            rows = []

            for values in data[1:]:
                row = {}

                for i, column in enumerate(columns):
                    value = values[i] if i < len(values) else ""
                    row[column] = "" if value is None else str(value)

                rows.append(row)

        # CSV
        else:
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                return {
                    "error": "Could not read the CSV encoding."
                }

            reader = csv.DictReader(io.StringIO(text))
            rows = list(reader)
            columns = reader.fieldnames or []

    except Exception as e:
        return {
            "error": f"Could not read the file: {str(e)}"
        }

    if not rows:
        return {
            "error": "CSV file is empty."
        }

    # -----------------------------
    # DATA QUALITY
    # -----------------------------

    missing_values = {}

    for column in columns:
        missing_values[column] = sum(
            1
            for row in rows
            if str(row.get(column, "")).strip() == ""
        )

    # Duplicate order IDs
    duplicate_orders = 0

    if "order_id" in columns:
        order_ids = [
            row.get("order_id", "").strip()
            for row in rows
            if row.get("order_id", "").strip()
        ]

        counts = Counter(order_ids)

        duplicate_orders = sum(
            count - 1
            for count in counts.values()
            if count > 1
        )

    # -----------------------------
    # BUSINESS CALCULATIONS
    # -----------------------------

    total_orders = len(rows)

    total_quantity = 0
    total_revenue = 0.0
    total_cost = 0.0
    returned_orders = 0

    product_revenue = defaultdict(float)
    category_revenue = defaultdict(float)
    region_revenue = defaultdict(float)
    monthly_revenue = defaultdict(float)

    product_orders = Counter()
    category_orders = Counter()
    region_orders = Counter()

    for row in rows:

        # Smart column detection
        quantity = to_int(
            row.get("quantity",
            row.get("qty",
            row.get("Qty", 0)))
        )

        unit_price = to_float(
            row.get("unit_price",
            row.get("price",
            row.get("Price", 0)))
        )

        cost_per_unit = to_float(
            row.get("cost_per_unit",
            row.get("cost",
            row.get("Cost", 0)))
        )

        discount = to_float(
            row.get("discount",
            row.get("Discount", 0))
        )

        # Revenue
        existing_revenue = to_float(
            row.get("revenue",
            row.get("Revenue", 0))
        )

        if existing_revenue > 0:
            revenue = existing_revenue
        else:
            gross_amount = quantity * unit_price
            discount_amount = gross_amount * (discount / 100)
            revenue = gross_amount - discount_amount

        # Cost
        if cost_per_unit > 0:
            cost = quantity * cost_per_unit
        else:
            cost = to_float(
                row.get("total_cost",
                row.get("Total Cost", 0))
            )

        # Profit
        profit = revenue - cost

        total_quantity += quantity
        total_revenue += revenue
        total_cost += cost

        # Return
        return_status = str(
            row.get("return_status", "")
        ).strip().lower()

        if return_status in ["yes", "returned", "true", "1"]:
            returned_orders += 1

        # Product
        product = row.get("product", "Unknown").strip()

        product_revenue[product] += revenue
        product_orders[product] += 1

        # Category
        category = row.get("category", "Unknown").strip()

        category_revenue[category] += revenue
        category_orders[category] += 1

        # Region
        region = row.get("region", "Unknown").strip()

        region_revenue[region] += revenue
        region_orders[region] += 1

        # Month
        order_date = row.get("order_date", "").strip()

        if order_date:

            try:
                date = datetime.strptime(
                    order_date,
                    "%Y-%m-%d"
                )

                month = date.strftime("%Y-%m")

                monthly_revenue[month] += revenue

            except:
                pass

    # -----------------------------
    # FINAL METRICS
    # -----------------------------

    total_profit = total_revenue - total_cost

    average_order_value = (
        total_revenue / total_orders
        if total_orders > 0
        else 0
    )

    return_rate = (
        (returned_orders / total_orders) * 100
        if total_orders > 0
        else 0
    )

    profit_margin = (
        (total_profit / total_revenue) * 100
        if total_revenue > 0
        else 0
    )

    # -----------------------------
    # TOP PERFORMERS
    # -----------------------------

    top_product = max(
        product_revenue,
        key=product_revenue.get
    ) if product_revenue else None

    top_category = max(
        category_revenue,
        key=category_revenue.get
    ) if category_revenue else None

    top_region = max(
        region_revenue,
        key=region_revenue.get
    ) if region_revenue else None

    # -----------------------------
    # SORTED REPORTS
    # -----------------------------

    top_products = sorted(
        [
            {
                "product": product,
                "revenue": round(revenue, 2),
                "orders": product_orders[product]
            }
            for product, revenue in product_revenue.items()
        ],
        key=lambda x: x["revenue"],
        reverse=True
    )

    top_categories = sorted(
        [
            {
                "category": category,
                "revenue": round(revenue, 2),
                "orders": category_orders[category]
            }
            for category, revenue in category_revenue.items()
        ],
        key=lambda x: x["revenue"],
        reverse=True
    )

    top_regions = sorted(
        [
            {
                "region": region,
                "revenue": round(revenue, 2),
                "orders": region_orders[region]
            }
            for region, revenue in region_revenue.items()
        ],
        key=lambda x: x["revenue"],
        reverse=True
    )

    monthly_report = {
        month: round(revenue, 2)
        for month, revenue in sorted(monthly_revenue.items())
    }

    # -----------------------------
    # AUTOMATIC INSIGHTS
    # -----------------------------

    insights = []

    if top_product:
        insights.append(
            f"{top_product} generated the highest product revenue."
        )

    if top_region:
        insights.append(
            f"{top_region} generated the highest regional revenue."
        )

    if return_rate > 10:
        insights.append(
            f"Return rate is {return_rate:.2f}%, which should be investigated."
        )

    if profit_margin < 15:
        insights.append(
            f"Profit margin is {profit_margin:.2f}%, indicating limited margin."
        )

    if duplicate_orders > 0:
        insights.append(
            f"{duplicate_orders} duplicate order records were detected."
        )

    if not insights:
        insights.append(
            "No major data-quality or performance warning was detected."
        )

    # -----------------------------
    # SAVE LATEST REPORT FOR EXCEL EXPORT
    # -----------------------------

    latest_report.update({
        "revenue": total_revenue,
        "cost": total_cost,
        "profit": total_profit,
        "orders": total_orders,
        "aov": average_order_value,
        "profit_margin": profit_margin,
        "return_rate": return_rate,
        "quantity_sold": total_quantity
    })

    # -----------------------------
    # FINAL RESPONSE
    # -----------------------------

    return {

        "status": "success",

        "file": {
            "filename": file.filename,
            "rows": total_orders,
            "columns": len(columns),
            "column_names": columns
        },

        "business_metrics": {

            "total_orders": total_orders,

            "total_quantity": total_quantity,

            "total_revenue": round(
                total_revenue,
                2
            ),

            "total_cost": round(
                total_cost,
                2
            ),

            "total_profit": round(
                total_profit,
                2
            ),

            "profit_margin_percent": round(
                profit_margin,
                2
            ),

            "average_order_value": round(
                average_order_value,
                2
            ),

            "returned_orders": returned_orders,

            "return_rate_percent": round(
                return_rate,
                2
            )
        },

        "top_performers": {

            "top_product": top_product,

            "top_category": top_category,

            "top_region": top_region
        },

        "product_analysis": top_products,

        "category_analysis": top_categories,

        "region_analysis": top_regions,

        "monthly_revenue": monthly_report,

        "data_quality": {

            "missing_values": missing_values,

            "duplicate_orders": duplicate_orders
        },

        "business_insights": insights
    }
from fastapi.responses import HTMLResponse
from pathlib import Path

@app.get("/dashboard", response_class=HTMLResponse)
def dashboard():
    html_path = Path(__file__).parent / "datapilot_dashboard.html"
    html = html_path.read_text(encoding="utf-8")
    return HTMLResponse(content=html)
@app.post("/export-excel")
async def export_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "DataPilot Report"

    ws.append(["DataPilot AI - Business Analysis Report"])
    ws.append([])
    ws.append(["Metric", "Value"])

    ws.append(["Revenue", latest_report.get("revenue", 0)])
    ws.append(["Cost", latest_report.get("cost", 0)])
    ws.append(["Profit", latest_report.get("profit", 0)])
    ws.append(["Orders", latest_report.get("orders", 0)])
    ws.append(["AOV", latest_report.get("aov", 0)])
    ws.append(["Profit Margin", latest_report.get("profit_margin", 0)])
    ws.append(["Return Rate", latest_report.get("return_rate", 0)])
    ws.append(["Quantity Sold", latest_report.get("quantity_sold", 0)])

    # Professional formatting
    from openpyxl.styles import Font, Alignment

    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A1:B1")

    ws["A3"].font = Font(bold=True)
    ws["B3"].font = Font(bold=True)

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 25

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition":
            'attachment; filename="DataPilot_AI_Business_Report.xlsx"'
        }
    )
